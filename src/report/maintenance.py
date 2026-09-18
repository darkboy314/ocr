"""Populate the maintenance tables in the 12-month report template."""

from __future__ import annotations

import csv
import re
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.table import Table, _Row
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DIRECTORY = PROJECT_ROOT / "input" / "Final_Algorithm"
DEFAULT_TEXT_DOCUMENT = DEFAULT_INPUT_DIRECTORY / "text.docx"
DEFAULT_GRAND_MASTER_CSV = DEFAULT_INPUT_DIRECTORY / "GrandMaster.csv"
DEFAULT_TEMPLATE = DEFAULT_INPUT_DIRECTORY / "12-month_Main_3_pairs_final.docx"
DEFAULT_RESULT_CSV = PROJECT_ROOT / "output" / "result.csv"
DEFAULT_OUTPUT_DOCUMENT = PROJECT_ROOT / "output" / "maintenance_report.docx"
DEFAULT_PROJECT_DICT = PROJECT_ROOT/ "input" / "Final_Algorithm" / "New_NETF_Trial_List_Master.csv"

DESCRIPTION_COLUMNS = tuple(f"main_desc_{index}" for index in range(1, 21))
COST_COLUMNS = tuple(f"main_cost_{index}" for index in range(1, 21))


def read_agreement_number(grand_master_csv: Path) -> str:
    """Return the agreement number from the single Generation 1 master-data row."""
    with grand_master_csv.open(newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.DictReader(csv_file)
        if reader.fieldnames is None:
            raise ValueError(f"The GrandMaster CSV is empty: {grand_master_csv}.")

        rows = [
            {
                header.strip(): _clean_csv_value(value or "")
                for header, value in row.items()
                if header is not None
            }
            for row in reader
        ]

    normalised_headers = {header.strip() for header in reader.fieldnames}
    required_columns = {"Generation", "Agr_Num"}
    missing_columns = sorted(required_columns.difference(normalised_headers))
    if missing_columns:
        raise ValueError(
            "The GrandMaster CSV is missing required columns: "
            f"{', '.join(missing_columns)}."
        )

    agreement_numbers = [
        row["Agr_Num"] for row in rows if row["Generation"].casefold() == "1"
    ]
    if len(agreement_numbers) != 1:
        raise ValueError(
            "Expected exactly one GrandMaster row with Generation = 1, found "
            f"{len(agreement_numbers)}."
        )
    if not agreement_numbers[0]:
        raise ValueError("The Generation = 1 GrandMaster row has an empty Agr_Num.")
    return agreement_numbers[0]


def read_maintenance_rows(csv_path: Path, agreement_number: str) -> list[dict[str, str]]:
    """Read rows for one agreement, normalising whitespace in CSV headers and values."""
    with csv_path.open(newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.reader(csv_file)
        try:
            headers = [header.strip() for header in next(reader)]
        except StopIteration as error:
            raise ValueError(f"The maintenance CSV is empty: {csv_path}.") from error

        required_columns = {
            "agreement_no",
            "veh_type",
            "license_plate",
            "total_main_type",
            "main_date_from",
            "main_date_to",
            "main_downtime",
            *DESCRIPTION_COLUMNS,
            *COST_COLUMNS,
        }
        missing_columns = sorted(required_columns.difference(headers))
        if missing_columns:
            raise ValueError(
                f"The maintenance CSV is missing required columns: {', '.join(missing_columns)}."
            )

        rows: list[dict[str, str]] = []
        agreement_column_index = headers.index("agreement_no")
        for values in reader:
            if (
                agreement_column_index >= len(values)
                or values[agreement_column_index].strip() != agreement_number
            ):
                continue
            values = _normalise_csv_values(values, len(headers))
            row = {
                header: _clean_csv_value(values[index]) if index < len(values) else ""
                for index, header in enumerate(headers)
            }
            rows.append(row)
    return rows


def _clean_csv_value(value: str) -> str:
    return value.strip().strip('"').strip()


def _normalise_csv_values(values: list[str], expected_count: int) -> list[str]:
    """Rejoin values split by commas inside a non-standard quoted CSV field."""
    normalised: list[str] = []
    index = 0
    while index < len(values):
        value = values[index]
        if value.lstrip().startswith('"') and not value.rstrip().endswith('"'):
            parts = [value]
            index += 1
            while index < len(values):
                parts.append(values[index])
                if values[index].rstrip().endswith('"'):
                    break
                index += 1
            normalised.append(",".join(parts))
        else:
            normalised.append(value)
        index += 1

    if len(normalised) != expected_count:
        raise ValueError(
            "A maintenance CSV row has an unexpected number of columns after "
            f"normalisation ({len(normalised)} instead of {expected_count})."
        )
    return normalised


def split_maintenance_rows(
    rows: Iterable[dict[str, str]],
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Separate scheduled and unscheduled records, discarding non-maintenance rows."""
    scheduled: list[dict[str, str]] = []
    unscheduled: list[dict[str, str]] = []
    for row in rows:
        maintenance_type = row["total_main_type"].casefold()
        if maintenance_type == "sch":
            scheduled.append(row)
        elif maintenance_type == "no_sch":
            unscheduled.append(row)
    return _ev_first(scheduled), _ev_first(unscheduled)


def _ev_first(rows: Iterable[dict[str, str]]) -> list[dict[str, str]]:
    return sorted(rows, key=lambda row: row["veh_type"].casefold() != "ev")


def _find_maintenance_table(document: Document, is_scheduled: bool) -> Table:
    for table in document.tables:
        header = " ".join(cell.text.casefold() for cell in table.rows[0].cells)
        has_maintenance_date = "maintenance" in header and "date" in header
        has_incident = "incident" in header
        if has_maintenance_date and has_incident != is_scheduled:
            return table

    table_name = "scheduled" if is_scheduled else "unscheduled"
    raise ValueError(f"Could not find the {table_name} maintenance table in the template.")


def _replace_rows(
    table: Table, records: list[dict[str, str]], has_incident_column: bool
) -> None:
    if len(table.rows) < 2:
        raise ValueError("The maintenance table has no data row to use as a formatting template.")

    row_template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)

    for record in records:
        new_row = _Row(deepcopy(row_template), table)
        for cell in new_row.cells:
            cell.text = ""
        _write_record(new_row, record, has_incident_column)
        table._tbl.append(new_row._tr)


def _write_record(
    row: _Row, record: dict[str, str], has_incident_column: bool
) -> None:
    expected_cells = 6
    if len(row.cells) < expected_cells:
        raise ValueError(
            f"A maintenance table row has fewer than {expected_cells} columns."
        )

    descriptions = _join_values(record, DESCRIPTION_COLUMNS)
    costs = _join_values(record, COST_COLUMNS)
    values = [
        _vehicle_label(record),
        *_unscheduled_leading_values(record, has_incident_column),
        _date_range(record),
        descriptions,
        record["main_downtime"],
        costs,
    ]
    if not has_incident_column:
        values.append(_total_cost(record))

    for cell, value in zip(row.cells, values, strict=True):
        _set_cell_text(cell, value)


def _join_values(record: dict[str, str], columns: Iterable[str]) -> str:
    return "\n".join(record[column] for column in columns if record[column])


def _unscheduled_leading_values(
    record: dict[str, str], has_incident_column: bool
) -> list[str]:
    return [record.get("accident_details", "")] if has_incident_column else []


def _total_cost(record: dict[str, str]) -> str:
    total_cost = record.get("total_cost", "")
    if total_cost:
        return total_cost

    costs = [record[column] for column in COST_COLUMNS if record[column]]
    if not costs:
        return ""

    amounts = [_parse_currency(cost) for cost in costs]
    total = sum(amounts, Decimal())
    return _format_currency(total)


def _parse_currency(value: str) -> Decimal:
    numeric_value = re.sub(r"[^\d.-]", "", value)
    try:
        return Decimal(numeric_value)
    except InvalidOperation as error:
        raise ValueError(f"Cannot calculate a total from maintenance cost {value!r}.") from error


def _format_currency(value: Decimal) -> str:
    if value == value.to_integral():
        return str(int(value))
    return format(value.normalize(), "f")


def _set_cell_text(cell, value: str) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _vehicle_label(record: dict[str, str]) -> str:
    vehicle_type = record["veh_type"].upper()
    license_plate = record["license_plate"]
    return f"{vehicle_type} ({license_plate})" if license_plate else vehicle_type


def _date_range(record: dict[str, str]) -> str:
    date_from = record["main_date_from"]
    date_to = record["main_date_to"]
    if date_from and date_to:
        return f"{date_from} – {date_to}"
    return date_from or date_to


def _clear_maintenance_caption_markers(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Table 3: Scheduled Maintenance"):
            paragraph.text = "Table 3: Scheduled Maintenance"
        elif paragraph.text.startswith("Table 4: Unscheduled Maintenance"):
            paragraph.text = "Table 4: Unscheduled Maintenance"


def _find_section_start(document: Document, heading: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(heading):
            return paragraph._p
    raise ValueError(f"Could not find section heading {heading!r}.")


def _replace_maintenance_section(source: Document, destination: Document) -> None:
    source_start = _find_section_start(source, "(b)")
    source_end = _find_section_start(source, "(c)")
    destination_start = _find_section_start(destination, "(b)")
    destination_end = _find_section_start(destination, "(c)")

    source_elements = list(source.element.body.iterchildren())
    destination_body = destination.element.body
    source_start_index = source_elements.index(source_start)
    source_end_index = source_elements.index(source_end)

    destination_elements = list(destination_body.iterchildren())
    destination_start_index = destination_elements.index(destination_start)
    destination_end_index = destination_elements.index(destination_end)

    for element in destination_elements[destination_start_index:destination_end_index]:
        destination_body.remove(element)

    insertion_index = destination_start_index
    for element in source_elements[source_start_index:source_end_index]:
        destination_body.insert(insertion_index, deepcopy(element))
        insertion_index += 1


def generate_maintenance_report(
    text_document: Path = DEFAULT_TEXT_DOCUMENT,
    template_document: Path = DEFAULT_TEMPLATE,
    csv_path: Path = DEFAULT_RESULT_CSV,
    output_document: Path = DEFAULT_OUTPUT_DOCUMENT,
    grand_master_csv: Path = DEFAULT_GRAND_MASTER_CSV,
) -> Path:
    """Create the maintenance report using agreement-matched data from ``result.csv``."""
    agreement_number = read_agreement_number(grand_master_csv)
    scheduled, unscheduled = split_maintenance_rows(
        read_maintenance_rows(csv_path, agreement_number)
    )

    template = Document(template_document)
    _replace_rows(
        _find_maintenance_table(template, is_scheduled=True),
        scheduled,
        has_incident_column=False,
    )
    _replace_rows(
        _find_maintenance_table(template, is_scheduled=False),
        unscheduled,
        has_incident_column=True,
    )
    _clear_maintenance_caption_markers(template)

    output_document.parent.mkdir(parents=True, exist_ok=True)
    template.save(output_document)

    source_report = Document(output_document)
    text_report = Document(text_document)
    _replace_maintenance_section(source_report, text_report)
    text_report.save(text_document)
    return output_document


if __name__ == "__main__":
    report_path = generate_maintenance_report()
    print(f"Created maintenance report: {report_path}")
